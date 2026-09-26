from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "submission" / "AURA_BUILD_LOG.docx"
FONT = "Arial"
NAVY = "17365D"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F5F6F7"
BORDER = "D9D9D9"


def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=115, bottom=85, end=115):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), BORDER)


def add_heading(document, value):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(1.5)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(value)
    set_run_font(run, size=11.5, bold=True, color="000000")


def add_body(document, value, bold_lead=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2.5)
    paragraph.paragraph_format.line_spacing = 1.03
    if bold_lead and value.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        set_run_font(first, size=10.3, bold=True)
        rest = paragraph.add_run(value[len(bold_lead):])
        set_run_font(rest, size=10.3)
    else:
        run = paragraph.add_run(value)
        set_run_font(run, size=10.3)
    return paragraph


def add_bullet(document, value):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(1.5)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(value)
    set_run_font(run, size=10.1)


document = Document()
section = document.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.05)
section.bottom_margin = Cm(1.0)
section.left_margin = Cm(1.35)
section.right_margin = Cm(1.35)

styles = document.styles
styles["Normal"].font.name = FONT
styles["Normal"].font.size = Pt(10.3)
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), FONT)
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

title = document.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(1)
title.paragraph_format.keep_with_next = True
# Override the blue bottom border carried by some Word Title style definitions.
title_ppr = title._p.get_or_add_pPr()
title_border = OxmlElement("w:pBdr")
title_bottom = OxmlElement("w:bottom")
title_bottom.set(qn("w:val"), "nil")
title_bottom.set(qn("w:sz"), "0")
title_bottom.set(qn("w:space"), "0")
title_border.append(title_bottom)
title_ppr.append(title_border)
run = title.add_run("AURA BUILD LOG SPRINT 1")
set_run_font(run, size=18, bold=True, color="000000")

subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(5)
run = subtitle.add_run("Track A The Escalation Referee | 27 09 2026 | Pham Gia Phu")
set_run_font(run, size=9.2, italic=True, color="555555")

intro = document.add_paragraph()
intro.paragraph_format.space_after = Pt(4)
intro.paragraph_format.line_spacing = 1.05
run = intro.add_run(
    "Trong Sprint 1, tôi dùng AI để rút ngắn thời gian đọc code, thử nghiệm kiến trúc và chuẩn hóa dữ liệu hóa đơn, "
    "nhưng giữ quyết định nghiệp vụ trong policy C# có thể kiểm thử. Kết quả là một vertical slice chạy thật từ upload, "
    "trích xuất và phân loại đến chuyển quản lý, audit và hoàn tác; các giới hạn của model được hiển thị thay vì che giấu."
)
set_run_font(run, size=10.4)

add_heading(document, "1 Công cụ và phạm vi sử dụng")
add_body(
    document,
    "Qwen3-VL-8B-Instruct qua OpenRouter đọc một ảnh JPG hoặc PNG và trả dữ kiện theo JSON Schema. "
    "Codex hỗ trợ truy route, phát hiện lỗi JavaScript, refactor, viết test và đồng bộ tài liệu. "
    "Tôi không dùng output của AI như bằng chứng duy nhất: thay đổi chỉ được giữ lại sau build, test và kiểm tra UI.",
)

add_heading(document, "2 Hành trình xây dựng và các lần AI sai")
add_body(
    document,
    "Prototype đầu tiên dùng Gemini Free Tier nhưng gặp HTTP 429 và quota không ổn định. Khi chuyển sang OpenRouter, "
    "một model slug hoặc endpoint sai từng gây HTTP 404; sau đó Qwen có lúc trả JSON lệch schema hoặc xem câu mô tả "
    "trên fixture tổng hợp là prompt injection. Ollama 4B còn từng đọc 295.199 VND thành 295.199 theo số thập phân "
    "và nhầm đơn giao hàng thành ride-hailing. Các lỗi này dẫn tới cấu hình provider rõ ràng, parser giới hạn, kiểm tra "
    "ngữ nghĩa và đúng một lượt đọc lại có mục tiêu; kết quả vẫn mâu thuẫn được chuyển cho người thay vì suy đoán.",
)

comparison = document.add_table(rows=1, cols=2)
comparison.alignment = WD_TABLE_ALIGNMENT.CENTER
comparison.autofit = False
comparison.columns[0].width = Cm(8.9)
comparison.columns[1].width = Cm(8.9)
set_table_borders(comparison)
for idx, label in enumerate(("AI giúp tăng tốc", "AI làm tốn công khi")):
    cell = comparison.rows[0].cells[idx]
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(label)
    set_run_font(r, size=9.5, bold=True, color="FFFFFF")

rows = [
    ("Tìm nhanh luồng upload, Verify, reviewer và audit.", "Đề xuất model hoặc API không tồn tại/chưa đúng slug."),
    ("Sinh test policy, contract và kiểm tra fixture offline.", "Output trông hợp lý nhưng không đúng schema hoặc policy."),
    ("Đối chiếu prompt, UI và tài liệu sau mỗi thay đổi.", "Tài liệu dễ tuyên bố quá mức nếu không đọc lại code thực tế."),
]
for row_idx, values in enumerate(rows):
    cells = comparison.add_row().cells
    for col_idx, value in enumerate(values):
        cell = cells[col_idx]
        set_cell_shading(cell, LIGHT_BLUE if row_idx % 2 == 0 else "FFFFFF")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_run_font(r, size=9.3)

add_heading(document, "3 Bằng chứng kiểm thử")
evidence = document.add_table(rows=2, cols=4)
evidence.alignment = WD_TABLE_ALIGNMENT.CENTER
evidence.autofit = False
widths = [Cm(4.35), Cm(4.35), Cm(4.35), Cm(4.35)]
headers = ("Build", "Test offline", "Verify v2 local", "Gói BGK")
values = ("0 warning 0 error", "70 trên 70 pass", "25 trên 25 fixture", "15 ca đã khóa")
for index, width in enumerate(widths):
    evidence.columns[index].width = width
set_table_borders(evidence)
for index, label in enumerate(headers):
    cell = evidence.rows[0].cells[index]
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(label)
    set_run_font(r, size=8.9, bold=True, color="FFFFFF")
for index, value in enumerate(values):
    cell = evidence.rows[1].cells[index]
    set_cell_shading(cell, LIGHT_GRAY if index % 2 else LIGHT_BLUE)
    set_cell_margins(cell, top=95, bottom=95)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(value)
    set_run_font(r, size=9.2, bold=True)

add_body(
    document,
    "Kết quả Ollama 25 trên 25 là năm batch liên tiếp của cùng 5 fixture tổng hợp đã biết trước; upload thủ công "
    "HoaDon1 đạt AUTO_APPROVE 3 trên 3. Các số liệu chứng minh pipeline, semantic repair và policy hoạt động nhất quán "
    "trên bộ kiểm soát, không phải độ chính xác trên hóa đơn thực tế độc lập.",
    bold_lead="Kết luận đo lường: ",
)

add_heading(document, "4 Cơ chế fail safe đã triển khai")
add_body(
    document,
    "AURA không retry HTTP 429 để tránh đốt thêm quota, chỉ retry tối đa một lần với lỗi 5xx và dừng các ca Verify còn "
    "lại khi provider gặp lỗi mang tính hệ thống. Đây chưa phải circuit breaker hoặc exponential backoff nhiều lần; "
    "mọi lỗi extraction đều chuyển kiểm tra thủ công và không tạo kết quả PASS giả.",
)

add_heading(document, "5 Quyết định kiến trúc và phần cắt giảm")
add_bullet(document, "Tách Qwen extraction và semantic validation khỏi PolicyDecisionEngine để đổi model mà không đổi quy tắc duyệt.")
add_bullet(document, "Giữ AuditLogs theo sự kiện append-only ở tầng ứng dụng, nhưng UI gom một hồ sơ thành một timeline để tránh cảm giác lặp.")
add_bullet(document, "Chặn thao tác workflow chồng trong một instance, dùng RowVersion chống ghi đè đồng thời và reload sau mutation để đọc lại trạng thái đã commit.")
add_bullet(document, "Giữ hosted Qwen 8B làm baseline; Ollama/Qwen 4B đã đạt 25/25 fixture nhưng vẫn tắt mặc định tới khi có benchmark OpenRouter cùng build và tập độc lập.")
add_bullet(document, "Hoãn PDF nhiều trang, authentication theo role, tax lookup, antivirus, object storage và benchmark tập dữ liệu độc lập.")

add_heading(document, "6 Bài học và bước tiếp theo")
add_body(
    document,
    "Lợi ích lớn nhất của AI là tăng tốc vòng lặp khám phá và kiểm thử, không phải thay người chịu trách nhiệm. "
    "AURA có đường chạy localhost tái lập trong README; bản SmarterASP.NET đã smoke thành công luồng upload, AI extraction "
    "và audit sau khi cập nhật API key trong Pool Manager. Gói Sprint 1 đã có video dưới ba phút, năm slide, Build Log, "
    "năm ca Verify và mười lăm ca tham chiếu cho BGK. Bước tiếp theo là pilot trên dữ liệu độc lập, đo field accuracy, "
    "over escalation, missed escalation và thời gian xử lý trước khi tuyên bố hiệu quả thực tế. "
    "Ảnh upload được gửi qua OpenRouter/provider Qwen; dữ liệu cá nhân thật không được đưa vào Git.",
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(OUTPUT)
