from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "submission" / "AURA_AFTER_FEEDBACK_REPORT.docx"

NAVY = "17324D"
TEAL = "1F8A8A"
PALE = "EFF6F8"
PALE_BLUE = "E8F0F6"
LIGHT = "F7F9FA"
GRAY = "D9D9D9"
INK = RGBColor(28, 35, 41)
MUTED = RGBColor(84, 96, 107)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRAY, size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_pr.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([run_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.75)
    section.footer_distance = Cm(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(8)

    for name, size in (("Heading 1", 17), ("Heading 2", 12.5), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "AURA   BÁO CÁO TRIỂN KHAI SAU FEEDBACK"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Arial"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("AURA   17 10 2026   •   ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)


def add_title(doc):
    p = doc.add_paragraph(style="Title")
    p.add_run("Báo cáo triển khai sau feedback doanh nghiệp")
    p_pr = p._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge_name in ("top", "left", "bottom", "right", "between", "bar"):
        edge = p_bdr.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            p_bdr.append(edge)
        edge.set(qn("w:val"), "nil")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("AURA ngày 17 tháng 10 năm 2026")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = MUTED

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Kết luận. ")
    r.bold = True
    p.add_run(
        "Candidate đã triển khai năm thay đổi trong cùng kiến trúc hiện hành gồm lượt chọn nhiều ảnh, "
        "số tiền AI suy ra có provenance, dashboard ngày tháng năm, báo cáo CSV và privacy triage fail safe. "
        "Release build đạt 0 warning 0 error, 181 trên 181 test offline. Thay đổi báo cáo không sửa database model."
    )
    p = doc.add_paragraph()
    p.add_run("Trạng thái live. ").bold = True
    p.add_run(
        "Baseline 5 trên 5 ngày 14 tháng 10 vẫn là bằng chứng của phiên bản trước feedback. "
        "Candidate mới chỉ được gọi là live sau publish, health check và synthetic smoke."
    )


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        if isinstance(item, tuple):
            lead, body = item
            p.add_run(lead).bold = True
            p.add_run(body)
        else:
            p.add_run(item)


def add_feedback_table(doc):
    doc.add_heading("Đối chiếu feedback và quyết định triển khai", level=1)
    rows = [
        ("Nhiều hóa đơn", "Làm ngay", "Tối đa 10 JPG hoặc PNG qua cùng HTTP 202 và SQL queue", "Mỗi ảnh là job riêng chưa có batchId hoặc PDF"),
        ("Folder hoặc email", "Roadmap", "Giữ ingestion boundary để nối connector sau", "Cần OIDC OAuth least privilege consent idempotency quarantine"),
        ("Không nhập số tiền", "Làm ngay", "UI bỏ ô amount; worker dùng final total AI sau validation", "Backend cũ chỉ giữ cho test kit; thiếu total thì fail safe"),
        ("Báo cáo quản lý", "Làm ngay", "Dashboard ngày tháng năm và CSV tháng hiện tại từ SQL", "Không xuất raw AI facts chưa phải financial ledger"),
        ("Dữ liệu cá nhân", "Làm ngay", "Category only triage và ESCALATE PRIVACY", "Không phải kết luận compliance"),
        ("Bảo vệ provider", "Làm ngay", "OpenRouter data collection deny", "Strict ZDR chờ xác nhận route live"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(3.0), Cm(2.2), Cm(7.4), Cm(4.8)]
    headers = ["Feedback", "Quyết định", "Đã triển khai", "Ranh giới"]
    for i, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = widths[i]
        cell.text = header
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
    repeat_header(table.rows[0])
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].width = widths[i]
            cells[i].text = value
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[i], top=55, bottom=55)
            if row_index % 2:
                set_cell_shading(cells[i], PALE_BLUE)
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(8.6)
        cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table)


def add_pipeline(doc):
    doc.add_page_break()
    doc.add_heading("Kiến trúc sau cập nhật", level=1)
    doc.add_paragraph(
        "Các boundary chính không đổi. Model chỉ đọc facts, C Sharp áp policy tất định và con người chịu trách nhiệm ngoại lệ. "
        "Privacy triage nằm sau structured extraction và semantic validation nhưng trước policy decision. "
        "Lượt chọn nhiều ảnh chỉ là tiện ích intake phía giao diện. Mỗi ảnh nhận status URL và SQL row riêng; "
        "candidate chưa có aggregate InvoiceBatch hoặc batch transaction phía server."
    )
    doc.add_paragraph(
        "Workflow mục tiêu giữ một ingestion boundary chung cho manual upload và mailbox connector. "
        "Mailbox Agent tương lai chỉ điều phối backend tool; connector riêng giữ OAuth và policy C Sharp vẫn là lớp duy nhất chọn route."
    )
    steps = [
        "Một đến mười JPG hoặc PNG cùng metadata; UI không nhận amount",
        "Kiểm extension MIME magic bytes và 5 MB mỗi ảnh",
        "Private storage SHA 256 SQL PENDING và audit AI QUEUED",
        "HTTP 202 và status URL riêng cho từng ảnh",
        "Worker claim bằng lease và coordinator giữ AI concurrency bằng một",
        "OpenRouter Qwen3.8 Flash reasoning none và data collection deny",
        "JSON Schema sang DTO rồi semantic validation và tối đa một repair",
        "Amount resolver và privacy classifier bổ sung provenance cùng category",
        "PolicyDecisionEngine trả AUTO APPROVE hoặc ESCALATE",
        "Nhân viên chuyển tiếp quản lý YES NO Undo audit dashboard và CSV tháng",
    ]
    for index, step in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.35)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(2.5)
        r = p.add_run(f"{index:02d}  ")
        r.bold = True
        r.font.color.rgb = RGBColor(31, 138, 138)
        p.add_run(step)


def add_safety(doc):
    doc.add_heading("Cơ chế an toàn mới", level=1)
    add_bullets(doc, [
        ("Lượt nhiều ảnh không làm mất job đã nhận. ", "UI lưu danh sách pending ngay sau từng HTTP 202 nên file sau lỗi không xóa SQL job trước; mỗi ảnh vẫn là một hồ sơ độc lập."),
        ("Không tin số tiền mù. ", "UI không gửi amount. Worker chỉ dùng final total AI lớn hơn không sau validation và audit nguồn; nhánh cũ được giữ cho test kit."),
        ("Privacy không quyết định bằng keyword. ", "Model phát marker category chuẩn backend chỉ nhận allow list và không ghi raw value vào warning. Vì triage chạy sau VLM demo ưu tiên chứng từ tổng hợp hoặc đã ẩn danh; production cần pre inference control cho dữ liệu nhạy cảm."),
        ("Fail safe có chủ đích. ", "Tín hiệu dữ liệu cá nhân dẫn đến ESCALATE PRIVACY fact inconsistency vẫn có ưu tiên cao hơn."),
        ("Provider boundary. ", "Payload yêu cầu data collection deny nhưng không bật strict ZDR nếu chưa xác nhận còn provider đủ điều kiện."),
        ("Dashboard và CSV dùng cùng SQL. ", "Không tạo source of truth mới, không xuất facts thô và vô hiệu hóa formula injection từ chuỗi người dùng."),
    ])
    p = doc.add_paragraph()
    p.add_run("Ranh giới pháp lý. ").bold = True
    p.add_run(
        "Luật Bảo vệ dữ liệu cá nhân 91 2025 QH15 và Nghị định 356 2025 NĐ CP có hiệu lực từ ngày 1 tháng 1 năm 2026. "
        "Tính năng này hỗ trợ sàng lọc vận hành chứ không thay thế đánh giá của pháp chế hoặc chứng nhận tuân thủ."
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    add_hyperlink(p, "Luật 91 2025 QH15", "https://vanban.chinhphu.vn/?docid=214590&pageid=27160")
    p.add_run("   •   ")
    add_hyperlink(p, "Nghị định 356 2025 NĐ CP", "https://vanban.chinhphu.vn/?docid=216387&pageid=27160&typegroupid=4")


def add_test_table(doc):
    doc.add_heading("Bằng chứng kiểm thử", level=1)
    data = [
        ("Release build", "PASS", "0 warning và 0 error"),
        ("Toàn bộ xUnit", "PASS", "181 trên 181"),
        ("Entity Framework", "PASS", "Không có pending model change"),
        ("Readiness pre check", "PASS", "0 failure hai warning đã biết"),
        ("Local health", "BLOCKED", "503 do LocalDB user instance của máy không khởi tạo được"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(5.0), Cm(2.4), Cm(10.0)]
    for i, header in enumerate(("Gate", "Kết quả", "Bằng chứng hoặc lưu ý")):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        cell.text = header
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    repeat_header(table.rows[0])
    for idx, row_data in enumerate(data):
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].width = widths[i]
            cells[i].text = value
            set_cell_margins(cells[i], top=35, bottom=35)
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if idx % 2:
                set_cell_shading(cells[i], PALE)
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(9)
        cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cells[1].paragraphs[0].runs[0].font.bold = True
        cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(31, 138, 138) if row_data[1] == "PASS" else RGBColor(171, 89, 26)
    set_table_borders(table)
    p = doc.add_paragraph()
    p.add_run("Diễn giải LocalDB. ").bold = True
    p.add_run(
        "Lỗi registry runtime trên máy kiểm tra không chứng minh database live hỏng và không phát sinh từ migration. "
        "Không xóa registry LocalDB hoặc MDF khi chưa có backup. Read only live health trả 200 nhưng dashboard candidate chưa xuất hiện và CSV trả 404 nên production vẫn là baseline cũ."
    )


def add_demo(doc):
    doc.add_heading("Quy trình demo mới", level=1)
    demo = [
        ("25 giây", "Mở Tổng quan", "Nêu số liệu từ SQL ở ba mức ngày tháng năm và tải CSV tháng."),
        ("35 giây", "Lượt nhiều ảnh", "Chọn hai hoặc ba ảnh; UI không có ô amount; chỉ HTTP 202 cùng job riêng."),
        ("25 giây", "Xử lý nền", "Chỉ PENDING hoặc PROCESSING và người dùng không bị giữ request."),
        ("25 giây", "Amount provenance", "Mở facts và chỉ AI EXTRACTED; thiếu total thì fail safe."),
        ("45 giây", "Privacy boundary", "Dùng synthetic receipt và chỉ ESCALATE PRIVACY không lặp raw value."),
        ("30 giây", "Human authority", "Nhân viên chuyển tiếp quản lý YES NO và AI không tự cấp quyền."),
        ("30 giây", "Audit và dashboard", "Mở timeline rồi refresh tổng quan."),
        ("30 giây", "Production reality", "Nêu data collection deny giới hạn ZDR và connector email còn ở roadmap."),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(2.3), Cm(4.0), Cm(11.1)]
    for i, header in enumerate(("Thời lượng", "Màn hình", "Ý phải nói")):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        cell.text = header
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    repeat_header(table.rows[0])
    for idx, row_data in enumerate(demo):
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].width = widths[i]
            cells[i].text = value
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if idx % 2:
                set_cell_shading(cells[i], PALE_BLUE)
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(8.2)
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table)
    p = doc.add_paragraph()
    p.add_run("Fallback khi provider chậm. ").bold = True
    p.add_run(
        "Chuyển sang audit đã ghi trong SQL hoặc video baseline. Không bấm Verify lặp lại để săn kết quả."
    )


def add_checklist_and_roadmap(doc):
    doc.add_heading("Checklist trước demo", level=1)
    checklist = [
        "Ghi commit SHA; backup hosting và database; kiểm biến môi trường mà không in API key.",
        "Publish và recycle rồi chờ 10 đến 30 giây.",
        "Kiểm health HTTP 200 status ok database storage true migration 0 và providerDataCollectionDenied true.",
        "Chạy single upload, xác nhận không có ô amount, lượt hai ảnh và một privacy synthetic case; kiểm manager route cùng audit.",
        "Không rerun holdout cũ để tạo accuracy claim mới; rehearsal theo thứ tự dashboard report multi file privacy manager audit.",
    ]
    for text in checklist:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.add_run("☐  ").font.color.rgb = RGBColor(31, 138, 138)
        p.add_run(text)

    doc.add_heading("Việc chưa triển khai và hướng tiếp theo", level=1)
    add_bullets(doc, [
        "Email hoặc folder ingestion dùng connector riêng giữ OAuth. Agent chỉ điều phối; cần scope tối thiểu allow list sync cursor idempotency malware scan quarantine webhook và backfill.",
        "Identity và dữ liệu: SSO RBAC segregation of duties private object storage KMS retention deletion và access logging.",
        "Multi instance: distributed limiter Verify lock circuit state cùng monitoring queue age error rate latency và capacity.",
        "Strict ZDR hoặc local VLM cho luồng nhạy cảm sau benchmark chất lượng latency và capacity.",
        "User study ba người theo protocol trước sau đã khóa ngưỡng và giữ phản hồi tiêu cực.",
    ])


def main():
    doc = Document()
    style_document(doc)
    add_title(doc)
    add_feedback_table(doc)
    add_pipeline(doc)
    add_safety(doc)
    add_test_table(doc)
    add_demo(doc)
    add_checklist_and_roadmap(doc)
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            keep_with_next(paragraph)
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Arial"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
            run.font.italic = False
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Arial"
                        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
                        run.font.italic = False

    doc.core_properties.title = "Báo cáo triển khai sau feedback doanh nghiệp"
    doc.core_properties.subject = "AURA implementation report 17 October 2026"
    doc.core_properties.author = "Phạm Gia Phú"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
