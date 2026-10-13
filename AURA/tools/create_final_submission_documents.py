from __future__ import annotations

import argparse
import html
import re
import sys
from pathlib import Path
from typing import Iterable, List, Sequence, Tuple

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    FrameBreak,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
SUBMISSION = ROOT / "submission"
TMP = ROOT / ".tmp" / "final-doc-build"
TMP.mkdir(parents=True, exist_ok=True)

BLUE = "17365D"
MID_BLUE = "2F5597"
LIGHT_BLUE = "D9EAF7"
PALE = "F3F6F9"
INK = "17212B"
MUTED = "5B6770"


def normalize_text(value: str) -> str:
    return (
        value.replace("→", "->")
        .replace("–", "-")
        .replace("—", "-")
        .replace("≤", "<=")
        .replace("≥", ">=")
        .replace("×", "x")
        .strip()
    )


def strip_inline_md(value: str) -> str:
    value = normalize_text(value)
    value = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", value)
    value = value.replace("**", "").replace("`", "")
    return value


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color: str = "CFD8E3", size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])


def apply_doc_defaults(doc: Document, compact: bool = False) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.62 if not compact else 0.44)
    section.bottom_margin = Inches(0.58 if not compact else 0.42)
    section.left_margin = Inches(0.70 if not compact else 0.48)
    section.right_margin = Inches(0.70 if not compact else 0.48)
    section.header_distance = Inches(0.24)
    section.footer_distance = Inches(0.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.2 if not compact else 8.35)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5 if not compact else 2.4)
    normal.paragraph_format.line_spacing = 1.08 if not compact else 1.0
    for name, size, color in (
        ("Title", 25 if not compact else 18, BLUE),
        ("Heading 1", 16 if not compact else 10.5, BLUE),
        ("Heading 2", 12.5 if not compact else 9.3, MID_BLUE),
        ("Heading 3", 10.8 if not compact else 8.7, INK),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8 if not compact else 3)
        style.paragraph_format.space_after = Pt(4 if not compact else 1.5)

    core = doc.core_properties
    core.author = ""
    core.last_modified_by = ""
    core.comments = ""
    core.company = ""
    core.keywords = ""

    header = section.header.paragraphs[0]
    header.text = "AURA  |  OrganizationAI  |  The Deciders"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(7.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(MID_BLUE)
    add_page_field(section.footer.paragraphs[0])


INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^\)]+\))")


def add_inline_runs(paragraph, value: str) -> None:
    value = normalize_text(value)
    for part in INLINE_RE.split(value):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor.from_string(MID_BLUE)
        elif part.startswith("["):
            match = re.match(r"\[([^\]]+)\]\([^\)]+\)", part)
            run = paragraph.add_run(match.group(1) if match else part)
            run.underline = True
            run.font.color.rgb = RGBColor.from_string(MID_BLUE)
        else:
            paragraph.add_run(part)


def parse_markdown(lines: Sequence[str]) -> List[Tuple[str, object]]:
    blocks: List[Tuple[str, object]] = []
    paragraph: List[str] = []
    i = 0

    def flush() -> None:
        if paragraph:
            blocks.append(("p", " ".join(x.strip() for x in paragraph)))
            paragraph.clear()

    while i < len(lines):
        raw = lines[i].rstrip("\n")
        line = raw.strip()
        if not line:
            flush()
            i += 1
            continue
        if line.startswith("```"):
            flush()
            code: List[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip("\n"))
                i += 1
            blocks.append(("code", "\n".join(code)))
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[i + 1].strip()):
            flush()
            rows: List[List[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(cells)
                i += 1
            if len(rows) >= 2:
                rows.pop(1)
            blocks.append(("table", rows))
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            flush()
            blocks.append((f"h{len(heading.group(1))}", heading.group(2)))
            i += 1
            continue
        if line.startswith("- "):
            flush()
            blocks.append(("bullet", line[2:].strip()))
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            flush()
            blocks.append(("number", re.sub(r"^\d+\.\s+", "", line)))
            i += 1
            continue
        if line.startswith(">"):
            flush()
            blocks.append(("quote", line.lstrip("> ")))
            i += 1
            continue
        paragraph.append(line)
        i += 1
    flush()
    return blocks


def add_docx_table(doc: Document, rows: Sequence[Sequence[str]], compact: bool = False) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for r_idx, row in enumerate(rows):
        for c_idx in range(width):
            cell = table.cell(r_idx, c_idx)
            set_cell_margins(cell, 45 if compact else 70, 60 if compact else 90, 45 if compact else 70, 60 if compact else 90)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx == 0:
                set_cell_shading(cell, BLUE)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, PALE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline_runs(paragraph, row[c_idx] if c_idx < len(row) else "")
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(7.7 if compact else 8.8)
                if r_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)


def add_blocks_to_docx(doc: Document, blocks: Sequence[Tuple[str, object]], compact: bool = False) -> None:
    first_h1 = True
    for kind, data in blocks:
        if kind == "h1":
            if first_h1:
                p = doc.add_paragraph(style="Title")
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_inline_runs(p, str(data))
                first_h1 = False
            else:
                add_inline_runs(doc.add_heading(level=1), str(data))
        elif kind == "h2":
            add_inline_runs(doc.add_heading(level=1), str(data))
        elif kind == "h3":
            add_inline_runs(doc.add_heading(level=2), str(data))
        elif kind == "p":
            p = doc.add_paragraph()
            add_inline_runs(p, str(data))
        elif kind in ("bullet", "number"):
            p = doc.add_paragraph(style="List Bullet" if kind == "bullet" else "List Number")
            p.paragraph_format.left_indent = Inches(0.16 if compact else 0.22)
            p.paragraph_format.first_line_indent = Inches(-0.12)
            add_inline_runs(p, str(data))
        elif kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.right_indent = Inches(0.14)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(6)
            add_inline_runs(p, str(data))
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor.from_string(MUTED)
        elif kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.16)
            p.paragraph_format.right_indent = Inches(0.16)
            set_cell = OxmlElement("w:shd")
            set_cell.set(qn("w:fill"), PALE)
            p._p.get_or_add_pPr().append(set_cell)
            run = p.add_run(str(data))
            run.font.name = "Consolas"
            run.font.size = Pt(7.5 if compact else 8.5)
        elif kind == "table":
            add_docx_table(doc, data, compact=compact)


def create_general_docx(source: Path, target: Path) -> None:
    doc = Document()
    apply_doc_defaults(doc)
    add_blocks_to_docx(doc, parse_markdown(source.read_text(encoding="utf-8").splitlines()))
    doc.save(target)


def create_build_log_docx(source: Path, target: Path) -> None:
    blocks = parse_markdown(source.read_text(encoding="utf-8").splitlines())
    doc = Document()
    apply_doc_defaults(doc, compact=True)
    title_seen = False
    split_index = 0
    for idx, (kind, data) in enumerate(blocks):
        if kind == "h1" and not title_seen:
            p = doc.add_paragraph(style="Title")
            add_inline_runs(p, str(data))
            title_seen = True
            continue
        if kind == "p" and str(data).startswith("**Track"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            add_inline_runs(p, str(data))
            split_index = idx + 1
            break
    section = doc.add_section(WD_SECTION.CONTINUOUS)
    section.top_margin = Inches(0.44)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.48)
    section.right_margin = Inches(0.48)
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")[0]
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "300")
    add_blocks_to_docx(doc, blocks[split_index:], compact=True)
    doc.save(target)


def add_form_field(doc: Document, label: str, lines: int = 2) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(label)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)
    for _ in range(lines):
        line = doc.add_paragraph("_" * 104)
        line.paragraph_format.space_after = Pt(1)
        for run in line.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string("A8B3BE")


def create_feedback_form_docx(target: Path) -> None:
    doc = Document()
    apply_doc_defaults(doc)
    title = doc.add_paragraph(style="Title")
    add_inline_runs(title, "AURA - Phiếu ghi nhận phản hồi người dùng")
    intro = doc.add_paragraph()
    add_inline_runs(intro, "Biểu mẫu trống cho ba phiên đánh giá thực tế. Dùng hóa đơn tổng hợp; không ghi API key, dữ liệu nội bộ hoặc thông tin nhạy cảm. Mỗi người tham gia phải đọc và xác nhận nguyên văn trước khi trích dẫn.")

    for idx in range(1, 4):
        if idx > 1:
            doc.add_page_break()
        doc.add_heading(f"Phiên người dùng {idx}", level=1)
        add_form_field(doc, "Họ tên thật:", 1)
        add_form_field(doc, "Chức danh và vai trò trong quy trình hoàn ứng:", 1)
        add_form_field(doc, "Ngày, giờ, địa điểm / hình thức thực hiện:", 1)
        p = doc.add_paragraph("[ ] Đồng ý tham gia tự nguyện    [ ] Đồng ý ghi âm/màn hình    [ ] Đồng ý công bố tên/chức danh    [ ] Chỉ cho phép tổng hợp ẩn danh")
        p.paragraph_format.space_after = Pt(5)

        doc.add_heading("Kết quả thực hiện tác vụ", level=2)
        rows = [
            ["Tác vụ", "Hoàn tất?", "Thời gian", "Vướng mắc / cần trợ giúp"],
            ["Upload ca thường và tìm kết quả", "", "", ""],
            ["Hiểu trạng thái ESCALATE và câu hỏi", "", "", ""],
            ["Chuyển tiếp, trả lời quản lý, xem Audit", "", "", ""],
            ["Undo quyết định gần nhất", "", "", ""],
        ]
        add_docx_table(doc, rows)
        add_form_field(doc, "Điểm hữu ích nhất:", 2)
        add_form_field(doc, "Điểm khó hiểu nhất:", 2)
        add_form_field(doc, "Một tác động tiêu cực/rủi ro cụ thể nếu dùng trong công việc:", 2)
        add_form_field(doc, "Trích dẫn nguyên văn do người tham gia tự viết hoặc đã đọc và xác nhận:", 3)
        add_form_field(doc, "Đề xuất ưu tiên số 1:", 2)
        add_form_field(doc, "Chữ ký/xác nhận của người tham gia và ngày:", 1)

    doc.add_page_break()
    doc.add_heading("Bảng tổng hợp minh bạch", level=1)
    add_docx_table(
        doc,
        [
            ["Người dùng", "Vai trò", "Trích dẫn được phép công bố", "Negative effect", "Ưu tiên"],
            ["1", "", "", "", ""],
            ["2", "", "", "", ""],
            ["3", "", "", "", ""],
        ],
    )
    add_form_field(doc, "Thay đổi sản phẩm đã chọn và lý do:", 3)
    add_form_field(doc, "Commit / before-after evidence:", 2)
    add_form_field(doc, "Điều gì không thay đổi và vì sao:", 2)
    add_form_field(doc, "Người tổng hợp, ngày xác nhận:", 1)
    doc.save(target)


def register_pdf_fonts() -> None:
    candidates = {
        "Arial": Path(r"C:\Windows\Fonts\arial.ttf"),
        "Arial-Bold": Path(r"C:\Windows\Fonts\arialbd.ttf"),
        "Arial-Italic": Path(r"C:\Windows\Fonts\ariali.ttf"),
        "Arial-Mono": Path(r"C:\Windows\Fonts\consola.ttf"),
    }
    for name, path in candidates.items():
        if path.exists():
            pdfmetrics.registerFont(TTFont(name, str(path)))
    pdfmetrics.registerFontFamily("Arial", normal="Arial", bold="Arial-Bold", italic="Arial-Italic", boldItalic="Arial-Bold")


def pdf_styles(compact: bool = False):
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "BodyAURA",
        parent=styles["BodyText"],
        fontName="Arial",
        fontSize=8.9 if not compact else 8.55,
        leading=11.2 if not compact else 10.45,
        textColor=colors.HexColor("#" + INK),
        spaceAfter=4 if not compact else 2,
    )
    return {
        "title": ParagraphStyle("TitleAURA", parent=body, fontName="Arial-Bold", fontSize=22 if not compact else 18, leading=25 if not compact else 20, textColor=colors.HexColor("#" + BLUE), spaceAfter=8 if not compact else 4),
        "h1": ParagraphStyle("H1AURA", parent=body, fontName="Arial-Bold", fontSize=14 if not compact else 10.6, leading=16 if not compact else 12, textColor=colors.HexColor("#" + BLUE), spaceBefore=8 if not compact else 4, spaceAfter=4 if not compact else 2, keepWithNext=True),
        "h2": ParagraphStyle("H2AURA", parent=body, fontName="Arial-Bold", fontSize=10.8 if not compact else 9.2, leading=13 if not compact else 10.6, textColor=colors.HexColor("#" + MID_BLUE), spaceBefore=6 if not compact else 3, spaceAfter=3 if not compact else 1.5, keepWithNext=True),
        "body": body,
        "bullet": ParagraphStyle("BulletAURA", parent=body, leftIndent=13, firstLineIndent=-7, bulletIndent=5, spaceAfter=2.5 if not compact else 1.4),
        "quote": ParagraphStyle("QuoteAURA", parent=body, leftIndent=14, rightIndent=8, textColor=colors.HexColor("#" + MUTED), fontName="Arial-Italic"),
        "code": ParagraphStyle("CodeAURA", parent=body, fontName="Arial-Mono", fontSize=7.2 if not compact else 6.8, leading=8.6, backColor=colors.HexColor("#" + PALE), leftIndent=8, rightIndent=8, borderPadding=5),
        "meta": ParagraphStyle("MetaAURA", parent=body, fontName="Arial-Bold", fontSize=8.6 if not compact else 8.3, textColor=colors.HexColor("#" + MUTED), spaceAfter=3),
    }


def md_inline_to_pdf(value: str) -> str:
    value = html.escape(normalize_text(value))
    value = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"<u><font color='#2F5597'>\1</font></u>", value)
    value = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", value)
    value = re.sub(r"`([^`]+)`", r"<font name='Arial-Mono' color='#2F5597'>\1</font>", value)
    return value


def make_pdf_table(rows: Sequence[Sequence[str]], available_width: float, compact: bool = False) -> Table:
    max_cols = max(len(row) for row in rows)
    if max_cols == 2:
        widths = [available_width * 0.64, available_width * 0.36]
    elif max_cols == 3:
        widths = [available_width * 0.30, available_width * 0.40, available_width * 0.30]
    elif max_cols == 4:
        widths = [available_width * 0.24, available_width * 0.18, available_width * 0.24, available_width * 0.34]
    elif max_cols == 5:
        widths = [available_width * 0.10, available_width * 0.15, available_width * 0.29, available_width * 0.27, available_width * 0.19]
    else:
        widths = [available_width / max_cols] * max_cols
    cell_style = ParagraphStyle("TableCell", fontName="Arial", fontSize=7.4 if compact else 7.5, leading=8.8 if compact else 9.1, textColor=colors.HexColor("#" + INK))
    head_style = ParagraphStyle("TableHead", parent=cell_style, fontName="Arial-Bold", textColor=colors.white)
    data = []
    for r_idx, row in enumerate(rows):
        data.append([Paragraph(md_inline_to_pdf(row[c] if c < len(row) else ""), head_style if r_idx == 0 else cell_style) for c in range(max_cols)])
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + BLUE)),
        ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#CFD8E3")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for row_idx in range(2, len(rows), 2):
        commands.append(("BACKGROUND", (0, row_idx), (-1, row_idx), colors.HexColor("#" + PALE)))
    table.setStyle(TableStyle(commands))
    return table


def header_footer(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("Arial-Bold", 7)
    canvas.setFillColor(colors.HexColor("#" + MID_BLUE))
    canvas.drawRightString(letter[0] - 0.65 * inch, letter[1] - 0.34 * inch, "AURA  |  OrganizationAI  |  The Deciders")
    canvas.setFont("Arial", 7)
    canvas.setFillColor(colors.HexColor("#" + MUTED))
    canvas.drawRightString(letter[0] - 0.65 * inch, 0.30 * inch, f"Page {doc.page}")
    canvas.restoreState()


def blocks_to_pdf_story(blocks: Sequence[Tuple[str, object]], width: float, compact: bool = False):
    styles = pdf_styles(compact)
    story = []
    first_h1 = True
    for kind, data in blocks:
        if kind == "h1":
            story.append(Paragraph(md_inline_to_pdf(str(data)), styles["title"] if first_h1 else styles["h1"]))
            first_h1 = False
        elif kind == "h2":
            story.append(Paragraph(md_inline_to_pdf(str(data)), styles["h1"]))
        elif kind == "h3":
            story.append(Paragraph(md_inline_to_pdf(str(data)), styles["h2"]))
        elif kind == "p":
            style = styles["meta"] if str(data).startswith("**Track") or str(data).startswith("**Đội") else styles["body"]
            story.append(Paragraph(md_inline_to_pdf(str(data)), style))
        elif kind in ("bullet", "number"):
            bullet = "•" if kind == "bullet" else "-"
            story.append(Paragraph(md_inline_to_pdf(str(data)), styles["bullet"], bulletText=bullet))
        elif kind == "quote":
            story.append(Paragraph(md_inline_to_pdf(str(data)), styles["quote"]))
        elif kind == "code":
            story.append(Paragraph(md_inline_to_pdf(str(data)).replace("\n", "<br/>"), styles["code"]))
        elif kind == "table":
            story.append(Spacer(1, 2))
            story.append(make_pdf_table(data, width, compact=compact))
            story.append(Spacer(1, 4))
    return story


def create_general_pdf(source: Path, target: Path) -> None:
    margin = 0.62 * inch
    doc = SimpleDocTemplate(str(target), pagesize=letter, leftMargin=margin, rightMargin=margin, topMargin=0.55 * inch, bottomMargin=0.48 * inch, title="AURA final submission", author="")
    width = letter[0] - 2 * margin
    blocks = parse_markdown(source.read_text(encoding="utf-8").splitlines())
    if source.name.startswith("AURA_PROPOSAL_"):
        before_privacy = []
        after_privacy = []
        destination = before_privacy
        for block in blocks:
            if block[0] == "h2" and str(block[1]).startswith("9 "):
                destination = after_privacy
            destination.append(block)
        story = blocks_to_pdf_story(before_privacy, width)
        story.append(PageBreak())
        story.extend(blocks_to_pdf_story(after_privacy, width))
    else:
        story = blocks_to_pdf_story(blocks, width)
    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)


def create_build_log_pdf(source: Path, target: Path) -> None:
    page_w, page_h = letter
    margin = 0.43 * inch
    header_height = 0.63 * inch
    footer_height = 0.28 * inch
    gap = 0.20 * inch
    usable_w = page_w - 2 * margin
    col_w = (usable_w - gap) / 2
    doc = BaseDocTemplate(str(target), pagesize=letter, leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin, title="AURA Build Log Sprint 2", author="")
    frames = [
        Frame(margin, page_h - margin - header_height, usable_w, header_height, id="header", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0),
        Frame(margin, margin + footer_height, col_w, page_h - 2 * margin - header_height - footer_height - 0.05 * inch, id="col1", leftPadding=0, rightPadding=5, topPadding=0, bottomPadding=0),
        Frame(margin + col_w + gap, margin + footer_height, col_w, page_h - 2 * margin - header_height - footer_height - 0.05 * inch, id="col2", leftPadding=5, rightPadding=0, topPadding=0, bottomPadding=0),
    ]
    doc.addPageTemplates(PageTemplate(id="buildlog", frames=frames, onPage=header_footer))
    blocks = parse_markdown(source.read_text(encoding="utf-8").splitlines())
    styles = pdf_styles(compact=True)
    story = []
    content_start = 0
    for idx, (kind, data) in enumerate(blocks):
        if kind == "h1":
            story.append(Paragraph(md_inline_to_pdf(str(data)), styles["title"]))
        elif kind == "p" and str(data).startswith("**Track"):
            story.append(Paragraph(md_inline_to_pdf(str(data)), styles["meta"]))
            story.append(FrameBreak())
            content_start = idx + 1
            break
    left_blocks = []
    right_blocks = []
    destination = left_blocks
    for block in blocks[content_start:]:
        if block[0] == "h2" and str(block[1]).startswith("What changed after failure"):
            destination = right_blocks
        destination.append(block)
    story.extend(blocks_to_pdf_story(left_blocks, col_w - 5, compact=True))
    story.append(FrameBreak())
    story.extend(blocks_to_pdf_story(right_blocks, col_w - 5, compact=True))
    doc.build(story)


def create_feedback_form_pdf(target: Path) -> None:
    margin = 0.58 * inch
    width = letter[0] - 2 * margin
    styles = pdf_styles()
    doc = SimpleDocTemplate(str(target), pagesize=letter, leftMargin=margin, rightMargin=margin, topMargin=0.54 * inch, bottomMargin=0.48 * inch, title="AURA user feedback form", author="")
    story = [
        Paragraph("AURA - Phiếu ghi nhận phản hồi người dùng", styles["title"]),
        Paragraph("Biểu mẫu trống cho ba phiên đánh giá thực tế. Dùng hóa đơn tổng hợp; không ghi API key, dữ liệu nội bộ hoặc thông tin nhạy cảm. Mỗi người tham gia phải đọc và xác nhận nguyên văn trước khi trích dẫn.", styles["body"]),
    ]

    def field(label: str, line_count: int = 2):
        result = [Paragraph(f"<b>{label}</b>", styles["body"])]
        for _ in range(line_count):
            result.append(Paragraph("________________________________________________________________________________________", ParagraphStyle("Line", parent=styles["body"], fontSize=7, textColor=colors.HexColor("#A8B3BE"), spaceAfter=1)))
        return result

    for idx in range(1, 4):
        if idx > 1:
            story.append(PageBreak())
        story.append(Paragraph(f"Phiên người dùng {idx}", styles["h1"]))
        story.extend(field("Họ tên thật:", 1))
        story.extend(field("Chức danh và vai trò trong quy trình hoàn ứng:", 1))
        story.extend(field("Ngày, giờ, địa điểm / hình thức thực hiện:", 1))
        story.append(Paragraph("[ ] Đồng ý tham gia tự nguyện &nbsp;&nbsp; [ ] Đồng ý ghi âm/màn hình &nbsp;&nbsp; [ ] Đồng ý công bố tên/chức danh &nbsp;&nbsp; [ ] Chỉ cho phép tổng hợp ẩn danh", styles["body"]))
        story.append(Paragraph("Kết quả thực hiện tác vụ", styles["h2"]))
        story.append(make_pdf_table([
            ["Tác vụ", "Hoàn tất?", "Thời gian", "Vướng mắc / cần trợ giúp"],
            ["Upload ca thường và tìm kết quả", "", "", ""],
            ["Hiểu trạng thái ESCALATE và câu hỏi", "", "", ""],
            ["Chuyển tiếp, trả lời quản lý, xem Audit", "", "", ""],
            ["Undo quyết định gần nhất", "", "", ""],
        ], width))
        story.extend(field("Điểm hữu ích nhất:", 2))
        story.extend(field("Điểm khó hiểu nhất:", 2))
        story.extend(field("Một tác động tiêu cực/rủi ro cụ thể nếu dùng trong công việc:", 2))
        story.extend(field("Trích dẫn nguyên văn do người tham gia tự viết hoặc đã đọc và xác nhận:", 3))
        story.extend(field("Đề xuất ưu tiên số 1:", 2))
        story.extend(field("Chữ ký/xác nhận của người tham gia và ngày:", 1))

    story.append(PageBreak())
    story.append(Paragraph("Bảng tổng hợp minh bạch", styles["h1"]))
    story.append(make_pdf_table([
        ["Người dùng", "Vai trò", "Trích dẫn được phép công bố", "Negative effect", "Ưu tiên"],
        ["1", "", "", "", ""],
        ["2", "", "", "", ""],
        ["3", "", "", "", ""],
    ], width))
    story.extend(field("Thay đổi sản phẩm đã chọn và lý do:", 3))
    story.extend(field("Commit / before-after evidence:", 2))
    story.extend(field("Điều gì không thay đổi và vì sao:", 2))
    story.extend(field("Người tổng hợp, ngày xác nhận:", 1))
    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--format", choices=("docx", "pdf", "all"), default="all")
    args = parser.parse_args()
    register_pdf_fonts()
    jobs = [
        ("AURA_PROPOSAL_VI_2026-10-13.md", "AURA_PROPOSAL_VI_2026-10-13"),
        ("AURA_PROPOSAL_EN_2026-10-13.md", "AURA_PROPOSAL_EN_2026-10-13"),
        ("AURA_SUBMISSION_LINKS_2026-10-13.md", "AURA_SUBMISSION_LINKS_2026-10-13"),
    ]
    for source_name, stem in jobs:
        source = SUBMISSION / source_name
        if args.format in ("docx", "all"):
            create_general_docx(source, TMP / f"{stem}.docx")
        if args.format in ("pdf", "all"):
            create_general_pdf(source, SUBMISSION / f"{stem}.pdf")
    build_source = SUBMISSION / "AURA_BUILD_LOG_1_PAGE_2026-10-13.md"
    if args.format in ("docx", "all"):
        create_build_log_docx(build_source, TMP / "AURA_BUILD_LOG_1_PAGE_2026-10-13.docx")
        create_feedback_form_docx(TMP / "AURA_USER_FEEDBACK_FORM_2026-10-13.docx")
    if args.format in ("pdf", "all"):
        create_build_log_pdf(build_source, SUBMISSION / "AURA_BUILD_LOG_1_PAGE_2026-10-13.pdf")
        create_feedback_form_pdf(SUBMISSION / "AURA_USER_FEEDBACK_FORM_2026-10-13.pdf")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
